"""Scrape the curated Tunisian shops into `ml/data/shops/local_catalogue.csv` (M7 local catalogue).

cd backend
python -m modules.m07_local_catalogue.scraper.run                      # all shops, 300 products max per shop
python -m modules.m07_local_catalogue.scraper.run --shops vog exist     # subset
python -m modules.m07_local_catalogue.scraper.run --max-per-shop 50 --no-xlsx --no-images

Pieces of foreign labels (Nike, Adidas, UGG...) are left out unless --keep-international is given.
Each product's photo is saved to ml/data/shops/images/ (column image_path); products without a photo
are dropped.
"""

import argparse
import csv
import logging
import random
import re
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

from ..paths import RAW_CSV
from .client import Http
from .images import download_images
from .normalize import categorize, fold, is_kids, is_non_fashion, join_sizes, main_color, origin_type
from .shops import SHOPS, Shop
from .sources import ADAPTERS, RawProduct

COLUMNS = [
    "product_id",
    "shop_name",
    "shop_url",
    "product_name",
    "category",
    "color",
    "size",
    "price_tnd",
    "city",
    "image_url",
    "image_path",
    "product_url",
    "available",
    "last_checked",
    "origin_type",
]
DEFAULT_OUT = RAW_CSV

log = logging.getLogger("localcatalogue.scraper")


def balanced_pick(
    rows: list[tuple[str, RawProduct]], cap: int, seed: int = 7
) -> list[tuple[str, RawProduct]]:
    """Round-robin over categories so a big shop does not flood the catalogue with one product type."""
    if len(rows) <= cap:
        return rows
    by_category: dict[str, list] = defaultdict(list)
    for row in rows:
        by_category[row[0]].append(row)
    rng = random.Random(seed)
    for bucket in by_category.values():
        rng.shuffle(bucket)
    picked: list[tuple[str, RawProduct]] = []
    while len(picked) < cap:
        for bucket in by_category.values():
            if bucket and len(picked) < cap:
                picked.append(bucket.pop())
    return picked


def scrape_shop(
    shop: Shop, http: Http, cap: int, include_kids: bool, keep_international: bool
) -> tuple[Shop, list[dict], dict]:
    cap = min(cap, shop.options.get("max_products", cap))
    stats = {
        "raw": 0,
        "kept": 0,
        "no_price": 0,
        "not_fashion": 0,
        "kids": 0,
        "international": 0,
        "duplicates": 0,
        "error": "",
    }
    try:
        raw = ADAPTERS[shop.platform](shop, http, max_raw=cap * 4)
    except Exception as exc:  # one broken shop must not stop the others
        log.exception("%s failed", shop.name)
        stats["error"] = f"{type(exc).__name__}: {exc}"[:200]
        return shop, [], stats
    stats["raw"] = len(raw)

    candidates: list[tuple[str, RawProduct]] = []
    seen: set = set()
    for p in raw:
        if not p.price_tnd or p.price_tnd <= 0 or not p.name or not p.product_url:
            stats["no_price"] += 1
            continue
        texts = [] if shop.options.get("name_only") else [*p.category_texts, p.tags]
        category = categorize(p.name, *texts) or shop.options.get("default_category")
        exclude = shop.options.get("exclude")
        if exclude and re.search(exclude, fold(" ".join([p.name, *p.category_texts]))):
            stats["not_fashion"] += 1
            continue
        if category is None or is_non_fashion(p.name):
            stats["not_fashion"] += 1
            continue
        if not include_kids and is_kids(p.name, *p.category_texts):
            stats["kids"] += 1
            continue
        # Filtered before the per-shop cap, so the quota is filled with local pieces.
        if not keep_international and origin_type(shop.origin, p.brand, p.name) == "international_brand":
            stats["international"] += 1
            continue
        # Same URL, or same name + price + photo (a piece listed under two categories).
        name_key = re.sub(r"\W+", "", p.name.lower())
        key = (p.product_url.rstrip("/").lower(), f"{name_key}|{p.price_tnd:.3f}|{p.image_url}")
        if key[0] in seen or key[1] in seen:
            stats["duplicates"] += 1
            continue
        seen.update(key)
        candidates.append((category, p))

    rows = []
    for category, p in balanced_pick(candidates, cap):
        if p.enrich:
            p.enrich(p)
        rows.append(
            {
                "shop_name": shop.name,
                "shop_url": shop.url.rstrip("/"),
                "product_name": p.name,
                "category": category,
                "color": main_color(p.name, p.colors, f"{p.tags} {p.extra_text}"),
                "size": join_sizes(p.sizes),
                "price_tnd": round(p.price_tnd, 3),
                "city": shop.city,
                "image_url": p.image_url.replace("http://", "https://")
                if p.image_url.startswith("http://")
                else ("https:" + p.image_url if p.image_url.startswith("//") else p.image_url),
                "product_url": p.product_url,
                "available": "true" if p.available else "false",
                "origin_type": origin_type(shop.origin, p.brand, p.name),
            }
        )
    stats["kept"] = len(rows)
    return shop, rows, stats


def write_outputs(rows: list[dict], out: Path, xlsx: bool) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so Excel shows accents (é, ç) correctly when the CSV is double-clicked.
    with out.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    if xlsx:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except ImportError:
            log.warning("openpyxl not installed: skipping the .xlsx copy")
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "local_catalogue"
        ws.append(COLUMNS)
        for row in rows:
            ws.append([row[c] for c in COLUMNS])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(i)].width = {
                "product_name": 45,
                "image_url": 40,
                "image_path": 26,
                "product_url": 40,
            }.get(col, 16)
        wb.save(out.with_suffix(".xlsx"))


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--shops", nargs="*", help="shop keys to scrape (default: all)")
    parser.add_argument("--max-per-shop", type=int, default=300)
    parser.add_argument(
        "--delay", type=float, default=0.5, help="seconds between two requests to the same site"
    )
    parser.add_argument(
        "--keep-international", action="store_true", help="keep pieces of foreign labels (Nike, Adidas...)"
    )
    parser.add_argument("--no-images", action="store_true", help="do not download the product photos")
    parser.add_argument("--include-kids", action="store_true", help="keep baby/children pieces")
    parser.add_argument("--no-xlsx", action="store_true", help="do not write the .xlsx copy")
    parser.add_argument(
        "--merge",
        action="store_true",
        help="with --shops: replace only these shops in the existing CSV and keep the others",
    )
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

    shops = [s for s in SHOPS if not args.shops or s.key in args.shops]
    http = Http(delay=args.delay)
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(
            pool.map(
                lambda s: scrape_shop(s, http, args.max_per_shop, args.include_kids, args.keep_international),
                shops,
            )
        )

    today = date.today().isoformat()
    for _, shop_rows, _ in results:
        for row in shop_rows:
            row["last_checked"] = today
    rows: list[dict] = []
    if args.merge and args.out.exists():
        scraped = {shop.name for shop in shops}
        with args.out.open(encoding="utf-8-sig") as fh:
            kept = [r for r in csv.DictReader(fh) if r["shop_name"] not in scraped]
        log.info("merge: keeping %d products of the other shops", len(kept))
        rows += kept
    for shop, shop_rows, stats in results:
        log.info(
            "%-18s raw=%-5d kept=%-4d no_price=%d not_fashion=%d kids=%d international=%d dup=%d %s",
            shop.name,
            stats["raw"],
            stats["kept"],
            stats["no_price"],
            stats["not_fashion"],
            stats["kids"],
            stats["international"],
            stats["duplicates"],
            stats["error"],
        )
        rows += sorted(shop_rows, key=lambda r: (r["category"], r["product_name"]))

    if args.no_images:
        for row in rows:
            row["image_path"] = ""
    else:
        log.info("downloading %d product photos...", len(rows))
        download_images(rows, Http(delay=0.2), args.out.parent / "images")
        missing = [row for row in rows if not row["image_path"]]
        log.info(
            "%d photos saved, %d products dropped (no usable photo)", len(rows) - len(missing), len(missing)
        )
        rows = [row for row in rows if row["image_path"]]
    for i, row in enumerate(rows, 1):
        row["product_id"] = f"P{i:04d}"

    write_outputs(rows, args.out, xlsx=not args.no_xlsx)
    log.info("wrote %d products from %d shops to %s", len(rows), sum(1 for _, r, _ in results if r), args.out)


if __name__ == "__main__":
    main()
